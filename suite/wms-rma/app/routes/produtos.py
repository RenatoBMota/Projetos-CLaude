import io
from flask import Blueprint, render_template, redirect, url_for, flash, request, send_file
from flask_login import login_required, current_user
from app.extensions import db
from app.models import Produto, Fornecedor

bp = Blueprint('produtos', __name__, url_prefix='/produtos')


@bp.route('/')
@login_required
def index():
    busca     = request.args.get('busca', '')
    forn_id   = request.args.get('fornecedor_id', type=int)
    categoria = request.args.get('categoria', '')
    page      = request.args.get('page', 1, type=int)

    q = Produto.query.filter_by(ativo=True)
    if busca:
        q = q.filter(db.or_(
            Produto.codigo.ilike(f'%{busca}%'),
            Produto.descricao.ilike(f'%{busca}%'),
            Produto.ean.ilike(f'%{busca}%'),
        ))
    if forn_id:
        q = q.filter_by(fornecedor_id=forn_id)
    if categoria:
        q = q.filter_by(categoria=categoria)

    produtos     = q.order_by(Produto.descricao).paginate(page=page, per_page=25, error_out=False)
    fornecedores = Fornecedor.query.filter_by(ativo=True).order_by(Fornecedor.nome).all()
    categorias   = db.session.query(Produto.categoria).filter(
        Produto.categoria.isnot(None)
    ).distinct().all()
    categorias   = [c[0] for c in categorias if c[0]]

    return render_template('produtos/index.html',
        produtos=produtos,
        fornecedores=fornecedores,
        categorias=categorias,
        filtro_busca=busca,
        filtro_forn=forn_id,
        filtro_categoria=categoria,
    )


@bp.route('/novo', methods=['GET', 'POST'])
@login_required
def novo():
    if request.method == 'POST':
        codigo = request.form.get('codigo', '').strip().upper()
        if Produto.query.filter_by(codigo=codigo).first():
            flash(f'Código "{codigo}" já cadastrado.', 'warning')
        else:
            p = Produto(
                codigo=codigo,
                descricao=request.form.get('descricao'),
                ean=request.form.get('ean'),
                marca=request.form.get('marca'),
                categoria=request.form.get('categoria'),
                fornecedor_id=request.form.get('fornecedor_id', type=int) or None,
                peso_kg=request.form.get('peso_kg', type=float),
                valor_unitario=request.form.get('valor_unitario', type=float),
                observacoes=request.form.get('observacoes'),
                comprador=request.form.get('comprador'),
            )
            db.session.add(p)
            db.session.commit()
            flash(f'Produto {codigo} cadastrado!', 'success')
            return redirect(url_for('produtos.index'))

    fornecedores = Fornecedor.query.filter_by(ativo=True).order_by(Fornecedor.nome).all()
    return render_template('produtos/form.html', produto=None, fornecedores=fornecedores)


@bp.route('/<int:prod_id>/editar', methods=['GET', 'POST'])
@login_required
def editar(prod_id):
    p = Produto.query.get_or_404(prod_id)
    if request.method == 'POST':
        p.descricao      = request.form.get('descricao')
        p.ean            = request.form.get('ean')
        p.marca          = request.form.get('marca')
        p.categoria      = request.form.get('categoria')
        p.fornecedor_id  = request.form.get('fornecedor_id', type=int) or None
        p.peso_kg        = request.form.get('peso_kg', type=float)
        p.valor_unitario = request.form.get('valor_unitario', type=float)
        p.observacoes    = request.form.get('observacoes')
        p.comprador      = request.form.get('comprador')
        db.session.commit()
        flash('Produto atualizado!', 'success')
        return redirect(url_for('produtos.index'))

    fornecedores = Fornecedor.query.filter_by(ativo=True).order_by(Fornecedor.nome).all()
    return render_template('produtos/form.html', produto=p, fornecedores=fornecedores)


@bp.route('/importar', methods=['GET', 'POST'])
@login_required
def importar():
    if request.method == 'POST':
        arquivo = request.files.get('arquivo')
        if not arquivo or not arquivo.filename.endswith(('.xls', '.xlsx')):
            flash('Envie um arquivo Excel (.xlsx ou .xls).', 'warning')
            return redirect(url_for('produtos.importar'))

        try:
            import openpyxl
            wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(min_row=2, values_only=True))
        except Exception as e:
            flash(f'Erro ao ler planilha: {e}', 'danger')
            return redirect(url_for('produtos.importar'))

        criados = atualizados = erros = 0
        fornecedores_cache = {f.nome.upper(): f.id for f in Fornecedor.query.all()}

        for i, row in enumerate(rows, start=2):
            if not row or not row[0]:
                continue
            try:
                codigo = str(row[0]).strip().upper()
                descricao = str(row[1]).strip() if row[1] else ''
                ean = str(row[2]).strip() if row[2] else None
                marca = str(row[3]).strip() if row[3] else None
                categoria = str(row[4]).strip() if row[4] else None
                forn_nome = str(row[5]).strip() if row[5] else None
                peso = float(row[6]) if row[6] else None
                valor = float(row[7]) if row[7] else None
                comprador = str(row[8]).strip() if len(row) > 8 and row[8] else None

                forn_id = None
                if forn_nome:
                    forn_id = fornecedores_cache.get(forn_nome.upper())

                p = Produto.query.filter_by(codigo=codigo).first()
                if p:
                    p.descricao = descricao or p.descricao
                    p.ean = ean; p.marca = marca; p.categoria = categoria
                    p.fornecedor_id = forn_id or p.fornecedor_id
                    p.peso_kg = peso; p.valor_unitario = valor; p.comprador = comprador
                    atualizados += 1
                else:
                    db.session.add(Produto(
                        codigo=codigo, descricao=descricao, ean=ean, marca=marca,
                        categoria=categoria, fornecedor_id=forn_id,
                        peso_kg=peso, valor_unitario=valor, comprador=comprador,
                    ))
                    criados += 1
            except Exception:
                erros += 1

        db.session.commit()
        flash(f'Importação concluída: {criados} criados, {atualizados} atualizados, {erros} erros.', 'success')
        return redirect(url_for('produtos.index'))

    return render_template('produtos/importar.html')


@bp.route('/template-excel')
@login_required
def template_excel():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Produtos'
    headers = ['codigo*', 'descricao*', 'ean', 'marca', 'categoria',
               'fornecedor_nome', 'peso_kg', 'valor_unitario', 'comprador']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(fill_type='solid', fgColor='2563EB')
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = 18
    # Example row
    ws.append(['PROD-001', 'Produto Exemplo', '7891234567890', 'Marca X',
               'Eletrodomésticos', 'Fornecedor ABC', 1.5, 299.90, 'João Silva'])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='template_produtos.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
