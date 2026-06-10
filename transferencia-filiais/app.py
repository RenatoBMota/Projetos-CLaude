import os
import io
import math
import sqlite3
from urllib.parse import quote, unquote
from flask import (Flask, render_template, request, redirect, url_for,
                   flash, send_file, abort, session)
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from db import init_db, get_sessions, get_session, create_session, insert_sugestoes
from db import get_sugestoes, get_sugestoes_by_comprador, get_compradores
from db import upsert_aprovacao, get_latest_session_id
from db import get_sugestao_by_codigo, get_all_compradores_in_session
from db import finalize_session, get_dashboard_data
from db import upsert_compradores_base, get_compradores_base, get_compradores_base_count
from calc import process_transfer

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'transferencia-filiais-secret-2024')
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB

UPLOAD_FOLDER = os.path.join(os.path.dirname(__file__), 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

ALLOWED_EXTENSIONS = {'csv', 'xlsx', 'xls'}


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def format_br(value, decimals=2):
    """Format number with Brazilian locale (. thousands, , decimals)."""
    if value is None:
        return '-'
    try:
        v = float(value)
        if v == 999.0:
            return '∞'
        formatted = f'{v:,.{decimals}f}'
        # Convert from US to BR format
        formatted = formatted.replace(',', 'X').replace('.', ',').replace('X', '.')
        return formatted
    except (TypeError, ValueError):
        return str(value)


def format_int_br(value):
    """Format integer with Brazilian locale."""
    if value is None:
        return '-'
    try:
        v = int(float(value))
        formatted = f'{v:,}'
        return formatted.replace(',', '.')
    except (TypeError, ValueError):
        return str(value)


app.jinja_env.filters['br'] = format_br
app.jinja_env.filters['br_int'] = format_int_br


@app.template_filter('url_encode')
def url_encode_filter(s):
    return quote(str(s), safe='')


@app.template_filter('decisao_badge')
def decisao_badge(decisao):
    if decisao == 'aprovado':
        return '<span class="badge bg-success">Aprovado</span>'
    elif decisao == 'recusado':
        return '<span class="badge bg-danger">Recusado</span>'
    elif decisao == 'alterado':
        return '<span class="badge bg-warning text-dark">Alterado</span>'
    else:
        return '<span class="badge bg-secondary">Pendente</span>'


@app.template_filter('status_badge')
def status_badge(status):
    mapping = {
        'Transferência Completa': 'bg-success',
        'Transferência Parcial': 'bg-warning text-dark',
        'Origem Insuficiente': 'bg-danger',
        'Sem Necessidade': 'bg-secondary',
    }
    cls = mapping.get(status, 'bg-secondary')
    return f'<span class="badge {cls}">{status}</span>'


with app.app_context():
    init_db()


# ─────────────────────────────────────────────
# HOME / UPLOAD
# ─────────────────────────────────────────────
@app.route('/', methods=['GET', 'POST'])
def index():
    sessions = get_sessions()

    base_info = get_compradores_base_count()

    if request.method == 'POST':
        # Validate required files
        required = ['vendas', 'estoque']
        for field in required:
            if field not in request.files or request.files[field].filename == '':
                flash(f'Arquivo "{field}" é obrigatório.', 'danger')
                return render_template('index.html', sessions=sessions, base_info=base_info)
            if not allowed_file(request.files[field].filename):
                flash(f'Formato inválido para "{field}". Use CSV ou XLSX.', 'danger')
                return render_template('index.html', sessions=sessions, base_info=base_info)

        # Read files into memory
        def get_file(field):
            f = request.files.get(field)
            if f and f.filename:
                data = f.read()
                return io.BytesIO(data), f.filename
            return None, None

        vendas_buf, vendas_name = get_file('vendas')
        estoque_buf, estoque_name = get_file('estoque')
        compradores_buf, compradores_name = get_file('compradores')
        transito_buf, transito_name = get_file('transito')
        reservas_buf, reservas_name = get_file('reservas')

        # Compradores: uploaded file takes priority; fall back to base cadastrada
        compradores_df_base = None
        if compradores_buf is None:
            base_rows = get_compradores_base()
            if not base_rows:
                flash('Nenhuma base de compradores cadastrada. Faça o upload do arquivo ou cadastre a base.', 'danger')
                return render_template('index.html', sessions=sessions, base_info=base_info)
            compradores_df_base = pd.DataFrame(base_rows)
            flash('Usando base de compradores cadastrada no sistema.', 'info')

        # Parameters
        try:
            filial_origem = request.form.get('filial_origem', '').strip().upper()
            filial_destino = request.form.get('filial_destino', '').strip().upper()
            dias_min_origem = int(request.form.get('dias_min_origem', 30))
            dias_meta_destino = int(request.form.get('dias_meta_destino', 15))
            nivel_servico = int(request.form.get('nivel_servico', 95))
        except ValueError as e:
            flash(f'Parâmetro inválido: {e}', 'danger')
            return render_template('index.html', sessions=sessions, base_info=base_info)

        if not filial_origem or not filial_destino:
            flash('Filial Origem e Filial Destino são obrigatórios.', 'danger')
            return render_template('index.html', sessions=sessions, base_info=base_info)

        try:
            results, warnings, meta = process_transfer(
                vendas_buf, vendas_name,
                estoque_buf, estoque_name,
                compradores_buf, compradores_name,
                transito_buf, transito_name,
                reservas_buf, reservas_name,
                filial_origem, filial_destino,
                dias_min_origem, dias_meta_destino, nivel_servico,
                compradores_df=compradores_df_base,
            )
        except ValueError as e:
            flash(str(e), 'danger')
            return render_template('index.html', sessions=sessions, base_info=base_info)
        except Exception as e:
            flash(f'Erro inesperado ao processar arquivos: {e}', 'danger')
            return render_template('index.html', sessions=sessions, base_info=base_info)

        for w in warnings:
            flash(w, 'warning')

        periodo_dias = meta['period_days']
        data_inicio  = meta['min_date']
        data_fim     = meta['max_date']
        flash(
            f'Período detectado: {data_inicio} a {data_fim} ({periodo_dias} dias). '
            f'Verifique se as datas do arquivo correspondem ao esperado.',
            'info'
        )

        total_produtos = sum(1 for r in results if r['incluido'])
        total_unidades = sum(r['sugestao'] for r in results if r['incluido'])

        session_id = create_session(
            filial_origem, filial_destino, periodo_dias, data_inicio, data_fim,
            dias_min_origem, dias_meta_destino, nivel_servico,
            total_produtos, total_unidades
        )
        insert_sugestoes(session_id, results)

        flash(f'Processamento concluído! {total_produtos} produtos encontrados.', 'success')
        return redirect(url_for('resultado', session_id=session_id))

    return render_template('index.html', sessions=sessions, base_info=base_info)


# ─────────────────────────────────────────────
# BASE DE COMPRADORES
# ─────────────────────────────────────────────
@app.route('/base-compradores', methods=['GET', 'POST'])
def base_compradores():
    base_info = get_compradores_base_count()

    if request.method == 'POST':
        f = request.files.get('compradores_base')
        if not f or not f.filename:
            flash('Selecione um arquivo para importar.', 'danger')
            return redirect(url_for('base_compradores'))
        if not allowed_file(f.filename):
            flash('Formato inválido. Use CSV ou XLSX.', 'danger')
            return redirect(url_for('base_compradores'))

        from calc import read_file, normalize_cod, find_col
        try:
            buf = io.BytesIO(f.read())
            df = read_file(buf, f.filename)
        except Exception as e:
            flash(f'Erro ao ler arquivo: {e}', 'danger')
            return redirect(url_for('base_compradores'))

        col_cod = find_col(df, 'codigo_produto')
        col_desc = find_col(df, 'descricao_produto')
        col_cod_forn = find_col(df, 'codigo_fornecedor')
        col_nome_forn = find_col(df, 'nome_fornecedor')
        col_comp = find_col(df, 'comprador')

        if not col_cod or not col_comp:
            flash('Arquivo deve conter colunas "Código Produto" e "Comprador".', 'danger')
            return redirect(url_for('base_compradores'))

        df[col_cod] = df[col_cod].apply(normalize_cod)
        rows = []
        for _, row in df.iterrows():
            cod = str(row[col_cod]).strip()
            if not cod or cod == 'nan':
                continue
            rows.append({
                'codigo_produto': cod,
                'descricao_produto': str(row[col_desc]).strip() if col_desc else cod,
                'codigo_fornecedor': str(row[col_cod_forn]).strip() if col_cod_forn else '',
                'nome_fornecedor': str(row[col_nome_forn]).strip() if col_nome_forn else '',
                'comprador': str(row[col_comp]).strip() if col_comp else 'Sem Comprador',
            })

        if not rows:
            flash('Nenhum registro válido encontrado no arquivo.', 'danger')
            return redirect(url_for('base_compradores'))

        upsert_compradores_base(rows)
        flash(f'Base atualizada com sucesso! {len(rows)} produtos cadastrados.', 'success')
        return redirect(url_for('base_compradores'))

    rows = get_compradores_base()
    compradores_unicos = sorted({r['comprador'] for r in rows})
    return render_template('base_compradores.html',
                           base_info=base_info,
                           rows=rows,
                           compradores_unicos=compradores_unicos)


# ─────────────────────────────────────────────
# RESULTADO
# ─────────────────────────────────────────────
@app.route('/resultado')
@app.route('/resultado/<int:session_id>')
def resultado(session_id=None):
    if session_id is None:
        session_id = get_latest_session_id()
        if session_id is None:
            flash('Nenhuma sessão encontrada. Faça um upload primeiro.', 'warning')
            return redirect(url_for('index'))
        return redirect(url_for('resultado', session_id=session_id))

    sess = get_session(session_id)
    if not sess:
        abort(404)

    rows = get_sugestoes(session_id)
    compradores_stats = get_compradores(session_id)

    # Filters
    filtro_status = request.args.get('status', '')
    filtro_comprador = request.args.get('comprador', '')
    filtro_busca = request.args.get('busca', '').lower()

    filtered = []
    for r in rows:
        if filtro_status and r['status'] != filtro_status:
            continue
        if filtro_comprador and r['comprador'] != filtro_comprador:
            continue
        if filtro_busca:
            if (filtro_busca not in str(r['codigo_produto']).lower() and
                    filtro_busca not in str(r['descricao_produto']).lower()):
                continue
        filtered.append(r)

    # Summary
    total = len(rows)
    total_units = sum(r['sugestao'] for r in rows)
    n_completo = sum(1 for r in rows if r['status'] == 'Transferência Completa')
    n_parcial = sum(1 for r in rows if r['status'] == 'Transferência Parcial')
    n_insuf = sum(1 for r in rows if r['status'] == 'Origem Insuficiente')

    # Unique statuses and compradores for filter dropdowns
    statuses = sorted(set(r['status'] for r in rows))
    compradores = sorted(set(r['comprador'] for r in rows))

    return render_template('resultado.html',
                           sess=sess,
                           rows=filtered,
                           total=total,
                           total_units=total_units,
                           n_completo=n_completo,
                           n_parcial=n_parcial,
                           n_insuf=n_insuf,
                           statuses=statuses,
                           compradores=compradores,
                           compradores_stats=compradores_stats,
                           filtro_status=filtro_status,
                           filtro_comprador=filtro_comprador,
                           filtro_busca=request.args.get('busca', ''))


# ─────────────────────────────────────────────
# EXPORT XLSX
# ─────────────────────────────────────────────
@app.route('/exportar/<int:session_id>')
def exportar(session_id):
    sess = get_session(session_id)
    if not sess:
        abort(404)

    rows = get_sugestoes(session_id)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Sugestões de Transferência'

    headers = [
        'Código', 'Produto', 'Comprador', 'Fornecedor',
        'MDV Destino', 'σ Desvio Padrão', 'CV', 'Est. Segurança',
        'Estoque Destino', 'Em Trânsito', 'Reservas',
        'Cobertura Atual (dias)', 'Estoque Desejado', 'Necessidade',
        'MDV Origem', 'Estoque Origem', 'Estoque Vital', 'Disponível',
        'Sugestão', 'Status', 'Decisão', 'Qtd Aprovada'
    ]

    # Header style
    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    center = Alignment(horizontal='center', vertical='center')
    thin = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    # Status colors
    fill_completo = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    fill_parcial = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
    fill_insuf = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')

    for row in rows:
        decisao = row['decisao'] or 'Pendente'
        qtd_aprov = row['quantidade_aprovada'] if row['quantidade_aprovada'] is not None else ''
        ws.append([
            row['codigo_produto'],
            row['descricao_produto'],
            row['comprador'],
            row['nome_fornecedor'],
            row['mdv_destino'],
            row['sigma_destino'],
            row['cv_destino'],
            row['estoque_seguranca'],
            row['estoque_destino'],
            row['em_transito'],
            row['reservas'],
            row['cobertura_destino_atual'],
            row['estoque_desejado'],
            row['necessidade'],
            row['mdv_origem'],
            row['estoque_origem'],
            row['estoque_vital'],
            row['disponivel'],
            row['sugestao'],
            row['status'],
            decisao.capitalize(),
            qtd_aprov,
        ])

        row_idx = ws.max_row
        status = row['status']
        if status == 'Transferência Completa':
            fill = fill_completo
        elif status == 'Transferência Parcial':
            fill = fill_parcial
        elif status == 'Origem Insuficiente':
            fill = fill_insuf
        else:
            fill = None

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border
            if fill:
                cell.fill = fill

    # Column widths
    col_widths = [12, 40, 20, 30, 12, 14, 12, 10, 18, 15, 12, 12, 14, 12, 12, 10, 22, 12, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = 'A2'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f'transferencia_{sess["filial_origem"]}_para_{sess["filial_destino"]}_{session_id}.xlsx'
    return send_file(output, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ─────────────────────────────────────────────
# COMPRADORES LIST
# ─────────────────────────────────────────────
@app.route('/compradores')
@app.route('/compradores/<int:session_id>')
def compradores(session_id=None):
    if session_id is None:
        session_id = get_latest_session_id()
        if session_id is None:
            flash('Nenhuma sessão encontrada.', 'warning')
            return redirect(url_for('index'))
        return redirect(url_for('compradores', session_id=session_id))

    sess = get_session(session_id)
    if not sess:
        abort(404)

    stats = get_compradores(session_id)
    return render_template('compradores.html', sess=sess, stats=stats)


# ─────────────────────────────────────────────
# APROVAÇÃO
# ─────────────────────────────────────────────
@app.route('/aprovacao/<comprador>', methods=['GET', 'POST'])
@app.route('/aprovacao/<comprador>/<int:session_id>', methods=['GET', 'POST'])
def aprovacao(comprador, session_id=None):
    comprador_decoded = unquote(comprador)

    if session_id is None:
        session_id = get_latest_session_id()
        if session_id is None:
            flash('Nenhuma sessão encontrada.', 'warning')
            return redirect(url_for('index'))
        return redirect(url_for('aprovacao', comprador=comprador, session_id=session_id))

    sess = get_session(session_id)
    if not sess:
        abort(404)

    rows = get_sugestoes_by_comprador(session_id, comprador_decoded)
    if not rows:
        flash(f'Comprador "{comprador_decoded}" não encontrado nesta sessão.', 'warning')
        return redirect(url_for('compradores', session_id=session_id))

    if request.method == 'POST':
        errors = []
        updates = []
        for row in rows:
            sid = str(row['id'])
            decisao = request.form.get(f'decisao_{sid}', '').strip()
            if not decisao:
                continue  # skip if no decision made
            if decisao not in ('aprovado', 'recusado', 'alterado'):
                errors.append(f'Decisão inválida para produto {row["codigo_produto"]}.')
                continue

            needs_qty = request.form.get(f'needs_qty_{sid}') == '1'

            if decisao == 'alterado' or (decisao == 'aprovado' and needs_qty):
                qtd_str = request.form.get(f'quantidade_{sid}', '').strip()
                try:
                    qtd = float(qtd_str.replace(',', '.'))
                    if qtd < 0:
                        raise ValueError
                except (ValueError, AttributeError):
                    errors.append(f'Quantidade inválida para produto {row["codigo_produto"]} (informe a quantidade).')
                    continue
                quantidade_aprovada = qtd
            elif decisao == 'aprovado':
                quantidade_aprovada = row['sugestao']
            else:  # recusado
                quantidade_aprovada = 0

            updates.append((row['id'], session_id, comprador_decoded, decisao, quantidade_aprovada))

        if errors:
            for e in errors:
                flash(e, 'danger')
        else:
            for args in updates:
                upsert_aprovacao(*args)
            flash(f'{len(updates)} decisão(ões) salva(s) com sucesso!', 'success')
            return redirect(url_for('aprovacao', comprador=comprador, session_id=session_id))

    # Summary
    n_total = len(rows)
    n_pendente = sum(1 for r in rows if r['decisao'] is None)
    n_aprovado = sum(1 for r in rows if r['decisao'] == 'aprovado')
    n_recusado = sum(1 for r in rows if r['decisao'] == 'recusado')
    n_alterado = sum(1 for r in rows if r['decisao'] == 'alterado')

    return render_template('aprovacao.html',
                           sess=sess,
                           comprador=comprador_decoded,
                           rows=rows,
                           n_total=n_total,
                           n_pendente=n_pendente,
                           n_aprovado=n_aprovado,
                           n_recusado=n_recusado,
                           n_alterado=n_alterado)


# ─────────────────────────────────────────────
# MEMÓRIA DE CÁLCULO / CONSULTA
# ─────────────────────────────────────────────
@app.route('/consulta/<int:session_id>', methods=['GET', 'POST'])
def consulta(session_id):
    sess = get_session(session_id)
    if not sess:
        abort(404)
    if request.method == 'POST':
        codigo = request.form.get('codigo', '').strip()
        if codigo:
            return redirect(url_for('consulta_produto', session_id=session_id, codigo=codigo))
    return render_template('consulta.html', sess=sess, produto=None, codigo_buscado='')


@app.route('/consulta/<int:session_id>/<path:codigo>')
def consulta_produto(session_id, codigo):
    from calc import Z_FACTORS
    sess = get_session(session_id)
    if not sess:
        abort(404)
    produto = get_sugestao_by_codigo(session_id, codigo)
    if not produto:
        flash(f'Produto "{codigo}" não encontrado nesta sessão.', 'warning')
    nivel_servico = sess['nivel_servico'] if sess['nivel_servico'] else 95
    z_factor = Z_FACTORS.get(int(nivel_servico), 1.65)
    return render_template('consulta.html', sess=sess, produto=produto,
                           codigo_buscado=codigo, z_factor=z_factor)


@app.route('/finalizar/<int:session_id>', methods=['POST'])
def finalizar(session_id):
    sess = get_session(session_id)
    if not sess:
        abort(404)
    if sess['finalizada']:
        flash('Sessão já estava finalizada.', 'warning')
    else:
        finalize_session(session_id)
        flash('Análise finalizada com sucesso. Nenhuma alteração adicional é permitida.', 'success')
    return redirect(url_for('resultado', session_id=session_id))


@app.route('/dashboard')
def dashboard():
    date_from = request.args.get('date_from', '')
    date_to = request.args.get('date_to', '')
    totals, buyers, sessions_list = get_dashboard_data(
        date_from or None,
        date_to or None
    )
    return render_template('dashboard.html',
        totals=totals, buyers=buyers, sessions_list=sessions_list,
        date_from=date_from, date_to=date_to)


if __name__ == '__main__':
    app.run(debug=True, port=5000)
