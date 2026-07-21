from flask import Blueprint, render_template, request, send_file, flash
from flask_login import login_required
from datetime import datetime, timedelta
from sqlalchemy import func
from app.extensions import db
from app.models import RMA, EstadoRMA, Fornecedor, Produto, PrazoSLA, Usuario
import io

bp = Blueprint('relatorios', __name__, url_prefix='/relatorios')


def _aplicar_filtros(q):
    """Aplica filtros comuns de data, estado, canal, fornecedor, departamento e loja."""
    data_ini     = request.args.get('data_ini') or request.form.get('data_ini')
    data_fim     = request.args.get('data_fim') or request.form.get('data_fim')
    estado       = request.args.get('estado')   or request.form.get('estado')
    canal        = request.args.get('canal')    or request.form.get('canal')
    forn_id      = request.args.get('fornecedor_id', type=int) or \
                   request.form.get('fornecedor_id', type=int)
    departamento = request.args.get('departamento') or request.form.get('departamento')
    loja_origem  = request.args.get('loja_origem') or request.form.get('loja_origem')

    filtros = {}
    if data_ini:
        dt = datetime.strptime(data_ini, '%Y-%m-%d')
        q = q.filter(RMA.criado_em >= dt)
        filtros['data_ini'] = data_ini
    if data_fim:
        dt = datetime.strptime(data_fim, '%Y-%m-%d') + timedelta(days=1)
        q = q.filter(RMA.criado_em < dt)
        filtros['data_fim'] = data_fim
    if estado:
        q = q.filter_by(estado=estado)
        filtros['estado'] = estado
    if canal:
        q = q.filter_by(canal=canal)
        filtros['canal'] = canal
    if forn_id:
        q = q.filter_by(fornecedor_id=forn_id)
        filtros['fornecedor_id'] = forn_id
    if departamento:
        q = q.join(Produto, RMA.produto_id == Produto.id)\
              .filter(Produto.categoria.ilike(f'%{departamento}%'))
        filtros['departamento'] = departamento
    if loja_origem:
        q = q.filter(RMA.loja_origem.ilike(f'%{loja_origem}%'))
        filtros['loja_origem'] = loja_origem

    return q, filtros


def _get_departamentos():
    rows = db.session.query(Produto.categoria)\
        .filter(Produto.categoria.isnot(None), Produto.categoria != '')\
        .distinct().order_by(Produto.categoria).all()
    return [r[0] for r in rows]


@bp.route('/')
@login_required
def index():
    fornecedores  = Fornecedor.query.filter_by(ativo=True).order_by(Fornecedor.nome).all()
    estados       = list(EstadoRMA.LABELS.items())
    departamentos = _get_departamentos()

    q, filtros = _aplicar_filtros(RMA.query)
    rmas = q.order_by(RMA.criado_em.desc()).all() if filtros else []

    total       = len(rmas)
    valor_total = sum(float(r.valor_total or 0) for r in rmas)
    finalizados = sum(1 for r in rmas if r.estado == EstadoRMA.FINALIZADO)
    atrasados   = sum(1 for r in rmas if r.em_atraso)

    por_estado = {}
    for r in rmas:
        por_estado[r.estado_label] = por_estado.get(r.estado_label, 0) + 1

    por_forn = {}
    for r in rmas:
        nome = r.fornecedor.nome if r.fornecedor else 'N/A'
        por_forn[nome] = por_forn.get(nome, 0) + 1

    aging = {'0-7 dias': 0, '8-15 dias': 0, '16-30 dias': 0, '31-60 dias': 0, '60+ dias': 0}
    for r in rmas:
        d = r.dias_em_aberto
        if d <= 7:    aging['0-7 dias'] += 1
        elif d <= 15: aging['8-15 dias'] += 1
        elif d <= 30: aging['16-30 dias'] += 1
        elif d <= 60: aging['31-60 dias'] += 1
        else:         aging['60+ dias'] += 1

    return render_template('relatorios/index.html',
        fornecedores=fornecedores,
        estados=estados,
        departamentos=departamentos,
        filtros=filtros,
        rmas=rmas,
        total=total,
        valor_total=valor_total,
        finalizados=finalizados,
        atrasados=atrasados,
        por_estado=por_estado,
        por_forn=por_forn,
        aging=aging,
    )


@bp.route('/exportar/excel')
@login_required
def exportar_excel():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        flash('openpyxl não instalado. Execute: pip install openpyxl', 'danger')
        return _redirect_relatorios()

    q, filtros = _aplicar_filtros(RMA.query)
    rmas = q.order_by(RMA.criado_em.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Relatório RMA'

    # Estilo do cabeçalho
    header_fill  = PatternFill('solid', fgColor='1E40AF')
    header_font  = Font(bold=True, color='FFFFFF', size=11)
    center_align = Alignment(horizontal='center', vertical='center')

    cabecalhos = [
        'Número', 'Estado', 'Canal', 'Loja Origem', 'Cliente',
        'Código Produto', 'EAN', 'Produto', 'Nº Série',
        'Fornecedor', 'Quantidade', 'NF Original', 'Motivo',
        'Categoria Defeito', 'Disposição', 'Valor Total (R$)',
        'Dias em Aberto', 'SLA em Atraso', 'Recebido Em', 'Finalizado Em',
    ]
    larguras = [18, 18, 10, 20, 25, 15, 15, 30, 18, 20, 10, 15, 30, 20, 20, 15, 12, 12, 18, 18]

    for col, (cab, larg) in enumerate(zip(cabecalhos, larguras), 1):
        cell = ws.cell(row=1, column=col, value=cab)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = center_align
        ws.column_dimensions[get_column_letter(col)].width = larg

    ws.row_dimensions[1].height = 22

    # Dados
    cor_par   = PatternFill('solid', fgColor='EFF6FF')
    cor_impar = PatternFill('solid', fgColor='FFFFFF')
    cor_atras = PatternFill('solid', fgColor='FEE2E2')

    for row_idx, rma in enumerate(rmas, 2):
        fill = cor_atras if rma.em_atraso else (cor_par if row_idx % 2 == 0 else cor_impar)
        dados = [
            rma.numero,
            rma.estado_label,
            rma.canal or '',
            rma.loja_origem or '',
            rma.cliente_nome or '',
            rma.produto.codigo if rma.produto else '',
            rma.produto.ean or '' if rma.produto else '',
            rma.produto.descricao if rma.produto else '',
            rma.numero_serie or '',
            rma.fornecedor.nome if rma.fornecedor else '',
            rma.quantidade,
            rma.nf_original or '',
            rma.motivo_devolucao or '',
            rma.categoria_defeito or '',
            rma.disposicao or '',
            float(rma.valor_total) if rma.valor_total else 0,
            rma.dias_em_aberto,
            'SIM' if rma.em_atraso else 'NÃO',
            rma.recebido_em.strftime('%d/%m/%Y %H:%M') if rma.recebido_em else '',
            rma.finalizado_em.strftime('%d/%m/%Y %H:%M') if rma.finalizado_em else '',
        ]
        for col, val in enumerate(dados, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill = fill

    # Totalizador no rodapé
    row_total = len(rmas) + 3
    ws.cell(row=row_total, column=1, value='TOTAL').font = Font(bold=True)
    ws.cell(row=row_total, column=11, value=sum(r.quantidade for r in rmas)).font = Font(bold=True)
    ws.cell(row=row_total, column=16, value=sum(float(r.valor_total or 0) for r in rmas)).font = Font(bold=True)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    nome = f'rma_relatorio_{datetime.utcnow().strftime("%Y%m%d_%H%M")}.xlsx'
    return send_file(output, as_attachment=True, download_name=nome,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@bp.route('/exportar/pdf')
@login_required
def exportar_pdf():
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    except ImportError:
        flash('reportlab não instalado. Execute: pip install reportlab', 'danger')
        return _redirect_relatorios()

    q, filtros = _aplicar_filtros(RMA.query)
    rmas = q.order_by(RMA.criado_em.desc()).all()

    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4),
                            rightMargin=1*cm, leftMargin=1*cm,
                            topMargin=1*cm, bottomMargin=1*cm)

    styles = getSampleStyleSheet()
    elementos = []

    # Título
    titulo_style = ParagraphStyle('titulo', parent=styles['Title'],
                                  fontSize=16, spaceAfter=6)
    sub_style = ParagraphStyle('sub', parent=styles['Normal'],
                               fontSize=9, textColor=colors.grey, spaceAfter=12)

    elementos.append(Paragraph('WMS RMA Enterprise — Relatório de RMAs', titulo_style))
    elementos.append(Paragraph(
        f'Gerado em {datetime.utcnow().strftime("%d/%m/%Y %H:%M")} | Total: {len(rmas)} registros',
        sub_style
    ))

    # Tabela
    cabecalhos_pdf = ['Número', 'Estado', 'Canal', 'Cliente', 'Produto',
                      'Fornecedor', 'Qtd', 'Dias', 'Atraso']
    dados_tabela = [cabecalhos_pdf]
    for rma in rmas:
        dados_tabela.append([
            rma.numero,
            rma.estado_label,
            rma.canal or '',
            (rma.cliente_nome or '')[:20],
            (rma.produto.descricao if rma.produto else '')[:25],
            (rma.fornecedor.nome if rma.fornecedor else '')[:18],
            str(rma.quantidade),
            str(rma.dias_em_aberto),
            '⚠ SIM' if rma.em_atraso else 'OK',
        ])

    estilo_tabela = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E40AF')),
        ('TEXTCOLOR',  (0, 0), (-1, 0), colors.white),
        ('FONTNAME',   (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE',   (0, 0), (-1, 0), 8),
        ('ALIGN',      (0, 0), (-1, 0), 'CENTER'),
        ('FONTSIZE',   (0, 1), (-1, -1), 7),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#EFF6FF')]),
        ('GRID',       (0, 0), (-1, -1), 0.3, colors.grey),
        ('VALIGN',     (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ])

    # Destacar linhas em atraso
    for i, rma in enumerate(rmas, 1):
        if rma.em_atraso:
            estilo_tabela.add('BACKGROUND', (0, i), (-1, i), colors.HexColor('#FEE2E2'))

    col_widths = [3.5*cm, 2.5*cm, 1.8*cm, 4*cm, 5*cm, 4*cm, 1*cm, 1.2*cm, 1.5*cm]
    tabela = Table(dados_tabela, colWidths=col_widths, repeatRows=1)
    tabela.setStyle(estilo_tabela)
    elementos.append(tabela)

    doc.build(elementos)
    output.seek(0)

    nome = f'rma_relatorio_{datetime.utcnow().strftime("%Y%m%d_%H%M")}.pdf'
    return send_file(output, as_attachment=True, download_name=nome, mimetype='application/pdf')


@bp.route('/aging')
@login_required
def aging():
    q = RMA.query.filter(RMA.estado.notin_([EstadoRMA.FINALIZADO, EstadoRMA.CANCELADO]))
    q, filtros = _aplicar_filtros(q)
    rmas_ativos = q.order_by(RMA.recebido_em).all()

    faixas = [
        ('0-7 dias',   0,   7),
        ('8-15 dias',  8,  15),
        ('16-30 dias', 16, 30),
        ('31-60 dias', 31, 60),
        ('60+ dias',   61, 9999),
    ]
    resultado = []
    for label, ini, fim in faixas:
        itens = [r for r in rmas_ativos if ini <= r.dias_em_aberto <= fim]
        valor = sum(float(r.valor_total or 0) for r in itens)
        resultado.append({'label': label, 'count': len(itens),
                          'valor': valor, 'rmas': itens[:10]})

    fornecedores  = Fornecedor.query.filter_by(ativo=True).order_by(Fornecedor.nome).all()
    departamentos = _get_departamentos()
    estados       = list(EstadoRMA.LABELS.items())
    return render_template('relatorios/aging.html', resultado=resultado, filtros=filtros,
        fornecedores=fornecedores, departamentos=departamentos, estados=estados)


@bp.route('/fornecedores')
@login_required
def por_fornecedor():
    q, filtros = _aplicar_filtros(RMA.query)

    dados = db.session.query(
        Fornecedor.nome,
        func.count(RMA.id).label('total'),
        func.sum(db.case((RMA.estado == EstadoRMA.FINALIZADO, 1), else_=0)).label('finalizados'),
        func.sum(db.case((RMA.estado == EstadoRMA.CANCELADO,  1), else_=0)).label('cancelados'),
        func.avg(func.julianday('now') - func.julianday(RMA.recebido_em)).label('media_dias'),
    ).join(Fornecedor, RMA.fornecedor_id == Fornecedor.id)\
     .group_by(Fornecedor.id)\
     .order_by(db.desc('total')).all()

    fornecedores  = Fornecedor.query.filter_by(ativo=True).order_by(Fornecedor.nome).all()
    departamentos = _get_departamentos()
    estados       = list(EstadoRMA.LABELS.items())
    return render_template('relatorios/fornecedores.html',
        dados=dados, fornecedores=fornecedores, filtros=filtros,
        departamentos=departamentos, estados=estados)


def _redirect_relatorios():
    from flask import redirect, url_for
    return redirect(url_for('relatorios.index'))
