import io
from flask import Blueprint, render_template, redirect, url_for, flash, request, send_file
from flask_login import login_required
from app.extensions import db
from app.models import Fornecedor

bp = Blueprint('fornecedores', __name__, url_prefix='/fornecedores')


@bp.route('/')
@login_required
def index():
    busca = request.args.get('busca', '')
    ativo = request.args.get('ativo', '')

    q = Fornecedor.query
    if busca:
        q = q.filter(db.or_(
            Fornecedor.nome.ilike(f'%{busca}%'),
            Fornecedor.cnpj.ilike(f'%{busca}%'),
            Fornecedor.contato.ilike(f'%{busca}%'),
        ))
    if ativo == '1':
        q = q.filter_by(ativo=True)
    elif ativo == '0':
        q = q.filter_by(ativo=False)
    else:
        q = q.filter_by(ativo=True)

    fornecedores = q.order_by(Fornecedor.nome).all()
    return render_template('fornecedores/index.html',
        fornecedores=fornecedores,
        filtro_busca=busca,
        filtro_ativo=ativo,
    )


@bp.route('/novo', methods=['GET', 'POST'])
@login_required
def novo():
    if request.method == 'POST':
        nome = request.form.get('nome', '').strip()
        cnpj = request.form.get('cnpj', '').strip()
        if not nome:
            flash('Nome é obrigatório.', 'warning')
        elif Fornecedor.query.filter_by(nome=nome).first():
            flash(f'Fornecedor "{nome}" já cadastrado.', 'warning')
        else:
            f = Fornecedor(
                nome=nome,
                cnpj=cnpj or None,
                email=request.form.get('email', '').strip() or None,
                telefone=request.form.get('telefone', '').strip() or None,
                contato=request.form.get('contato', '').strip() or None,
            )
            db.session.add(f)
            db.session.commit()
            flash(f'Fornecedor "{nome}" cadastrado!', 'success')
            return redirect(url_for('fornecedores.index'))

    return render_template('fornecedores/form.html', fornecedor=None)


@bp.route('/<int:forn_id>/editar', methods=['GET', 'POST'])
@login_required
def editar(forn_id):
    f = Fornecedor.query.get_or_404(forn_id)
    if request.method == 'POST':
        f.nome     = request.form.get('nome', f.nome).strip()
        f.cnpj     = request.form.get('cnpj', '').strip() or None
        f.email    = request.form.get('email', '').strip() or None
        f.telefone = request.form.get('telefone', '').strip() or None
        f.contato  = request.form.get('contato', '').strip() or None
        f.ativo    = request.form.get('ativo') == 'on'
        db.session.commit()
        flash('Fornecedor atualizado!', 'success')
        return redirect(url_for('fornecedores.index'))

    return render_template('fornecedores/form.html', fornecedor=f)


@bp.route('/importar', methods=['GET', 'POST'])
@login_required
def importar():
    if request.method == 'POST':
        arquivo = request.files.get('arquivo')
        if not arquivo or not arquivo.filename.endswith(('.xls', '.xlsx')):
            flash('Envie um arquivo Excel (.xlsx ou .xls).', 'warning')
            return redirect(url_for('fornecedores.importar'))

        try:
            import openpyxl
            wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(min_row=2, values_only=True))
        except Exception as e:
            flash(f'Erro ao ler planilha: {e}', 'danger')
            return redirect(url_for('fornecedores.importar'))

        criados = atualizados = ignorados = 0
        for i, row in enumerate(rows, start=2):
            if not row or not row[0]:
                continue
            try:
                nome    = str(row[0]).strip() if row[0] else None
                cnpj    = str(row[1]).strip() if len(row) > 1 and row[1] else None
                email   = str(row[2]).strip() if len(row) > 2 and row[2] else None
                telefone= str(row[3]).strip() if len(row) > 3 and row[3] else None
                contato = str(row[4]).strip() if len(row) > 4 and row[4] else None

                if not nome:
                    ignorados += 1
                    continue

                f = Fornecedor.query.filter(
                    db.func.lower(Fornecedor.nome) == nome.lower()
                ).first()

                if f:
                    if cnpj:    f.cnpj     = cnpj
                    if email:   f.email    = email
                    if telefone:f.telefone = telefone
                    if contato: f.contato  = contato
                    atualizados += 1
                else:
                    db.session.add(Fornecedor(
                        nome=nome, cnpj=cnpj, email=email,
                        telefone=telefone, contato=contato,
                    ))
                    criados += 1
            except Exception:
                ignorados += 1

        db.session.commit()
        flash(f'Importação concluída: {criados} criados, {atualizados} atualizados, {ignorados} ignorados.', 'success')
        return redirect(url_for('fornecedores.index'))

    return render_template('fornecedores/importar.html')


@bp.route('/template-excel')
@login_required
def template_excel():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Fornecedores'
    headers = ['nome*', 'cnpj', 'email', 'telefone', 'contato']
    larguras = [40, 20, 30, 18, 30]
    for col, (h, w) in enumerate(zip(headers, larguras), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(fill_type='solid', fgColor='2563EB')
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = w
    ws.append(['Bosch Brasil', '07.144.457/0001-65', 'rma@bosch.com.br', '(11) 4444-0001', 'Maria Silva'])
    ws.append(['Samsung Brasil', '24.030.072/0001-98', 'rma@samsung.com.br', '(11) 4444-0002', 'João Costa'])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='template_fornecedores.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
