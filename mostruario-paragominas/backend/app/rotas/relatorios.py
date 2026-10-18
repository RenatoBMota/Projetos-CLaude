"""Indicadores e relatórios.

Todos respeitam a abrangência de filiais do usuário. Os relatórios em tabela
podem ser baixados em Excel (formato=xlsx) ou CSV (formato=csv).
"""
import csv
import io
from dataclasses import dataclass
from datetime import date, timedelta

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from psycopg import Connection

from ..db import get_conn
from ..seguranca import exigir_perfil, filiais_permitidas

router = APIRouter(prefix="/api/relatorios", tags=["relatorios"])

pode_ver_relatorios = exigir_perfil("gestor", "admin", "aprovador_comercial", "aprovador_logistica", "deposito")

# Movimentações válidas = não são estorno e não foram estornadas.
MOV_VALIDAS = """
    SELECT m.*, i.solicitacao_id, i.produto_id, s.filial_id, s.protocolo, s.requisitante_id
    FROM movimentacoes m
    JOIN itens_solicitacao i ON i.id = m.item_id
    JOIN solicitacoes s ON s.id = i.solicitacao_id
    WHERE m.tipo <> 'estorno'
      AND NOT EXISTS (SELECT 1 FROM movimentacoes e WHERE e.estorno_de_id = m.id)
"""


@dataclass
class Filtro:
    inicio: date
    fim: date
    filial_id: int | None
    fornecedor_id: int | None
    filiais: list[int] | None  # abrangência; None = todas

    def params(self) -> dict:
        # fim inclusivo: compara com < fim + 1 dia
        return {"inicio": self.inicio, "fim_exclusivo": self.fim + timedelta(days=1),
                "filial_id": self.filial_id, "fornecedor_id": self.fornecedor_id, "filiais": self.filiais}

    def filial(self, coluna: str) -> str:
        """Condição SQL de filial (abrangência + filtro escolhido)."""
        partes = []
        if self.filiais is not None:
            partes.append(f"{coluna} = ANY(%(filiais)s)")
        if self.filial_id:
            partes.append(f"{coluna} = %(filial_id)s")
        return " AND ".join(partes) or "true"

    def fornecedor(self, coluna: str) -> str:
        return f"{coluna} = %(fornecedor_id)s" if self.fornecedor_id else "true"


def filtro(inicio: date | None = None, fim: date | None = None, filial_id: int | None = None,
           fornecedor_id: int | None = None, usuario: dict = Depends(pode_ver_relatorios)) -> Filtro:
    fim = fim or date.today()
    inicio = inicio or (fim - timedelta(days=89))
    if inicio > fim:
        raise HTTPException(400, "A data inicial é maior que a final")
    permitidas = filiais_permitidas(usuario)
    return Filtro(inicio, fim, filial_id, fornecedor_id, list(permitidas) if permitidas is not None else None)


def _razao(a, b):
    return round(float(a) / float(b), 4) if b else None


# ---------------------------------------------------------------------------
# Indicadores
# ---------------------------------------------------------------------------

@router.get("/indicadores")
def indicadores(f: Filtro = Depends(filtro), conn: Connection = Depends(get_conn)):
    p = f.params()
    um = lambda sql: conn.execute(sql, p).fetchone()  # noqa: E731

    req = um(f"""
        SELECT count(DISTINCT s.id) AS solicitacoes,
               COALESCE(sum(i.quantidade_solicitada), 0) AS qtd_solicitada,
               COALESCE(sum(i.quantidade_solicitada * i.custo_referencia), 0) AS valor_solicitado
        FROM solicitacoes s JOIN itens_solicitacao i ON i.solicitacao_id = s.id
        JOIN produtos p ON p.id = i.produto_id
        WHERE s.status <> 'rascunho' AND s.enviado_em >= %(inicio)s AND s.enviado_em < %(fim_exclusivo)s
          AND {f.filial('s.filial_id')} AND {f.fornecedor('p.fornecedor_id')}""")

    mov = um(f"""
        SELECT
          COALESCE(sum(m.quantidade) FILTER (WHERE m.tipo = 'expedicao'), 0) AS qtd_expedida,
          COALESCE(sum(m.quantidade * m.custo_unitario) FILTER (WHERE m.tipo = 'expedicao'), 0) AS valor_expedido,
          COALESCE(sum(m.quantidade) FILTER (WHERE m.tipo = 'devolucao_estoque'), 0) AS qtd_estoque,
          COALESCE(sum(m.quantidade * m.custo_unitario) FILTER (WHERE m.tipo = 'devolucao_estoque'), 0) AS valor_estoque,
          COALESCE(sum(m.quantidade) FILTER (WHERE m.tipo = 'devolucao_descarte'), 0) AS qtd_descarte,
          COALESCE(sum(m.quantidade * m.custo_unitario) FILTER (WHERE m.tipo = 'devolucao_descarte'), 0) AS valor_descarte,
          COALESCE(sum(m.quantidade) FILTER (WHERE m.tipo = 'perda'), 0) AS qtd_perda,
          COALESCE(sum(m.quantidade * m.custo_unitario) FILTER (WHERE m.tipo = 'perda'), 0) AS valor_perda
        FROM ({MOV_VALIDAS}) m JOIN produtos p ON p.id = m.produto_id
        WHERE m.ocorrido_em >= %(inicio)s AND m.ocorrido_em < %(fim_exclusivo)s
          AND {f.filial('m.filial_id')} AND {f.fornecedor('p.fornecedor_id')}""")

    expo = um(f"""
        SELECT COALESCE(sum(v.qtd_em_loja), 0) AS qtd, COALESCE(sum(v.valor_em_loja), 0) AS valor,
               count(*) AS itens, round(avg(v.dias_em_exposicao), 1) AS idade_media
        FROM vw_posicao_item v
        WHERE v.qtd_em_loja > 0 AND {f.filial('v.filial_id')} AND {f.fornecedor('v.fornecedor_id')}""")

    permanencia = um(f"""
        SELECT round(avg(v.dias_em_exposicao), 1) AS dias
        FROM vw_posicao_item v JOIN solicitacoes s ON s.id = v.solicitacao_id
        WHERE v.status_item = 'encerrado' AND s.encerrado_em >= %(inicio)s AND s.encerrado_em < %(fim_exclusivo)s
          AND {f.filial('v.filial_id')} AND {f.fornecedor('v.fornecedor_id')}""")

    aprov = conn.execute(f"""
        SELECT a.etapa, count(*) FILTER (WHERE a.decisao = 'aprovada') AS aprovadas, count(*) AS decididas
        FROM aprovacoes a JOIN solicitacoes s ON s.id = a.solicitacao_id
        WHERE a.decidido_em >= %(inicio)s AND a.decidido_em < %(fim_exclusivo)s AND {f.filial('s.filial_id')}
        GROUP BY a.etapa""", p).fetchall()
    taxas = {a["etapa"]: _razao(a["aprovadas"], a["decididas"]) for a in aprov}

    # Tempos médios em horas
    tempos = um(f"""
        WITH t AS (
          SELECT s.id, s.enviado_em,
                 (SELECT decidido_em FROM aprovacoes WHERE solicitacao_id = s.id AND etapa = 'comercial') AS com,
                 (SELECT decidido_em FROM aprovacoes WHERE solicitacao_id = s.id AND etapa = 'logistica'
                    AND decisao = 'aprovada') AS log,
                 (SELECT min(m.ocorrido_em) FROM movimentacoes m JOIN itens_solicitacao i ON i.id = m.item_id
                   WHERE i.solicitacao_id = s.id AND m.tipo = 'expedicao') AS exp
          FROM solicitacoes s
          WHERE s.enviado_em >= %(inicio)s AND s.enviado_em < %(fim_exclusivo)s AND {f.filial('s.filial_id')}
        )
        SELECT round(avg(extract(epoch FROM com - enviado_em)) / 3600, 1) AS horas_comercial,
               round(avg(extract(epoch FROM log - com)) / 3600, 1) AS horas_logistica,
               round(avg(extract(epoch FROM exp - log)) / 3600, 1) AS horas_expedicao
        FROM t""")

    pend = um(f"""
        SELECT count(*) FILTER (WHERE status IN ('aguardando_aprovacao_comercial', 'aguardando_aprovacao_logistica')) AS aprovacao,
               count(*) FILTER (WHERE status = 'aguardando_expedicao') AS expedicao
        FROM solicitacoes s WHERE {f.filial('s.filial_id')}""")

    retornado = mov["qtd_estoque"] + mov["qtd_descarte"] + mov["qtd_perda"]
    return {
        "periodo": {"inicio": f.inicio, "fim": f.fim},
        "solicitacoes": req["solicitacoes"],
        "qtd_solicitada": req["qtd_solicitada"],
        "valor_solicitado": req["valor_solicitado"],
        **mov,
        "em_exposicao": {"quantidade": expo["qtd"], "valor": expo["valor"], "itens": expo["itens"],
                         "idade_media_dias": expo["idade_media"]},
        "permanencia_media_dias": permanencia["dias"],
        "taxa_aprovacao_comercial": taxas.get("comercial"),
        "taxa_aprovacao_logistica": taxas.get("logistica"),
        "horas_aprovacao_comercial": tempos["horas_comercial"],
        "horas_aprovacao_logistica": tempos["horas_logistica"],
        "horas_ate_expedicao": tempos["horas_expedicao"],
        # bases: o que retornou ou foi encerrado no período
        "taxa_reaproveitamento": _razao(mov["qtd_estoque"], retornado),
        "taxa_descarte": _razao(mov["qtd_descarte"], retornado),
        "taxa_perda": _razao(mov["qtd_perda"], retornado),
        "pendentes_aprovacao": pend["aprovacao"],
        "pendentes_expedicao": pend["expedicao"],
    }


# ---------------------------------------------------------------------------
# Relatórios em tabela
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class Relatorio:
    titulo: str
    colunas: list[tuple[str, str, str]]  # (campo, rótulo, tipo: texto|qtd|moeda|data|datahora|pct|int)
    sql: str  # usa {filial(col)}, {fornecedor(col)} e {MOV}
    usa_periodo: bool = True


EXPOSICAO_AGRUPADA = """
    SELECT {chave}, count(*) AS itens, sum(v.qtd_em_loja) AS quantidade, sum(v.valor_em_loja) AS valor,
           round(avg(v.dias_em_exposicao), 1) AS idade_media
    FROM vw_posicao_item v {juncao}
    WHERE v.qtd_em_loja > 0 AND {{filial(v.filial_id)}} AND {{fornecedor(v.fornecedor_id)}}
    GROUP BY {grupo} ORDER BY valor DESC
"""

COLS_EXPOSICAO = [("itens", "Itens", "int"), ("quantidade", "Quantidade", "qtd"),
                  ("valor", "Valor em exposição", "moeda"), ("idade_media", "Idade média (dias)", "qtd")]

RELATORIOS: dict[str, Relatorio] = {
    "exposicao_filial": Relatorio(
        "Em exposição por filial", [("filial", "Filial", "texto")] + COLS_EXPOSICAO,
        EXPOSICAO_AGRUPADA.format(chave="f.nome AS filial", juncao="JOIN filiais f ON f.id = v.filial_id",
                                  grupo="f.nome"), usa_periodo=False),
    "exposicao_fornecedor": Relatorio(
        "Em exposição por fornecedor", [("fornecedor", "Fornecedor", "texto")] + COLS_EXPOSICAO,
        EXPOSICAO_AGRUPADA.format(chave="COALESCE(fo.razao_social, '(sem fornecedor)') AS fornecedor",
                                  juncao="LEFT JOIN fornecedores fo ON fo.id = v.fornecedor_id",
                                  grupo="fo.razao_social"), usa_periodo=False),
    "exposicao_produto": Relatorio(
        "Em exposição por produto",
        [("sku", "SKU", "texto"), ("descricao", "Produto", "texto")] + COLS_EXPOSICAO,
        EXPOSICAO_AGRUPADA.format(chave="p.sku, p.descricao", juncao="JOIN produtos p ON p.id = v.produto_id",
                                  grupo="p.sku, p.descricao"), usa_periodo=False),
    "exposicao_requisitante": Relatorio(
        "Em exposição por requisitante", [("requisitante", "Requisitante", "texto")] + COLS_EXPOSICAO,
        EXPOSICAO_AGRUPADA.format(chave="u.nome AS requisitante", juncao="JOIN usuarios u ON u.id = v.requisitante_id",
                                  grupo="u.nome"), usa_periodo=False),
    "movimentacoes": Relatorio(
        "Movimentações no período",
        [("ocorrido_em", "Data", "datahora"), ("protocolo", "Protocolo", "texto"), ("filial", "Filial", "texto"),
         ("tipo", "Tipo", "texto"), ("sku", "SKU", "texto"), ("descricao", "Produto", "texto"),
         ("fornecedor", "Fornecedor", "texto"), ("quantidade", "Quantidade", "qtd"), ("valor", "Valor", "moeda"),
         ("responsavel", "Responsável", "texto"), ("motivo", "Motivo", "texto"), ("observacao", "Observação", "texto")],
        """SELECT m.ocorrido_em, m.protocolo, f.nome AS filial,
                  CASE m.tipo WHEN 'expedicao' THEN 'Expedição' WHEN 'devolucao_estoque' THEN 'Devolução → Estoque'
                       WHEN 'devolucao_descarte' THEN 'Devolução → Descarte' ELSE 'Perda' END AS tipo,
                  p.sku, p.descricao, fo.razao_social AS fornecedor, m.quantidade,
                  m.quantidade * m.custo_unitario AS valor, u.nome AS responsavel, mo.descricao AS motivo, m.observacao
           FROM ({MOV}) m JOIN filiais f ON f.id = m.filial_id JOIN produtos p ON p.id = m.produto_id
           LEFT JOIN fornecedores fo ON fo.id = p.fornecedor_id JOIN usuarios u ON u.id = m.usuario_id
           LEFT JOIN motivos mo ON mo.id = m.motivo_id
           WHERE m.ocorrido_em >= %(inicio)s AND m.ocorrido_em < %(fim_exclusivo)s
             AND {filial(m.filial_id)} AND {fornecedor(p.fornecedor_id)}
           ORDER BY m.ocorrido_em DESC LIMIT 5000"""),
    "descarte_perda": Relatorio(
        "Descarte e perda por filial",
        [("filial", "Filial", "texto"), ("qtd_estoque", "Voltou p/ estoque", "qtd"),
         ("qtd_descarte", "Descarte", "qtd"), ("valor_descarte", "Valor descartado", "moeda"),
         ("qtd_perda", "Perda", "qtd"), ("valor_perda", "Valor perdido", "moeda"),
         ("taxa_descarte", "% descarte", "pct"), ("taxa_perda", "% perda", "pct")],
        """SELECT f.nome AS filial,
                  sum(m.quantidade) FILTER (WHERE m.tipo = 'devolucao_estoque') AS qtd_estoque,
                  sum(m.quantidade) FILTER (WHERE m.tipo = 'devolucao_descarte') AS qtd_descarte,
                  sum(m.quantidade * m.custo_unitario) FILTER (WHERE m.tipo = 'devolucao_descarte') AS valor_descarte,
                  sum(m.quantidade) FILTER (WHERE m.tipo = 'perda') AS qtd_perda,
                  sum(m.quantidade * m.custo_unitario) FILTER (WHERE m.tipo = 'perda') AS valor_perda,
                  round(COALESCE(sum(m.quantidade) FILTER (WHERE m.tipo = 'devolucao_descarte'), 0) / sum(m.quantidade), 4) AS taxa_descarte,
                  round(COALESCE(sum(m.quantidade) FILTER (WHERE m.tipo = 'perda'), 0) / sum(m.quantidade), 4) AS taxa_perda
           FROM ({MOV}) m JOIN filiais f ON f.id = m.filial_id JOIN produtos p ON p.id = m.produto_id
           WHERE m.tipo <> 'expedicao' AND m.ocorrido_em >= %(inicio)s AND m.ocorrido_em < %(fim_exclusivo)s
             AND {filial(m.filial_id)} AND {fornecedor(p.fornecedor_id)}
           GROUP BY f.nome ORDER BY valor_perda DESC NULLS LAST, valor_descarte DESC NULLS LAST"""),
    "motivos": Relatorio(
        "Descarte e perda por motivo",
        [("destino", "Destino", "texto"), ("motivo", "Motivo", "texto"), ("ocorrencias", "Ocorrências", "int"),
         ("quantidade", "Quantidade", "qtd"), ("valor", "Valor", "moeda")],
        """SELECT CASE m.tipo WHEN 'perda' THEN 'Perda' ELSE 'Descarte' END AS destino, mo.descricao AS motivo,
                  count(*) AS ocorrencias, sum(m.quantidade) AS quantidade, sum(m.quantidade * m.custo_unitario) AS valor
           FROM ({MOV}) m JOIN motivos mo ON mo.id = m.motivo_id JOIN produtos p ON p.id = m.produto_id
           WHERE m.tipo IN ('devolucao_descarte', 'perda')
             AND m.ocorrido_em >= %(inicio)s AND m.ocorrido_em < %(fim_exclusivo)s
             AND {filial(m.filial_id)} AND {fornecedor(p.fornecedor_id)}
           GROUP BY 1, 2 ORDER BY valor DESC"""),
    "lancar_erp": Relatorio(
        "Devoluções para estoque (lançar no ERP)",
        [("sku", "SKU", "texto"), ("descricao", "Produto", "texto"), ("unidade", "Unidade", "texto"),
         ("quantidade", "Quantidade", "qtd"), ("valor", "Valor", "moeda"), ("protocolos", "Protocolos", "texto")],
        """SELECT p.sku, p.descricao, p.unidade, sum(m.quantidade) AS quantidade,
                  sum(m.quantidade * m.custo_unitario) AS valor,
                  string_agg(DISTINCT m.protocolo, ', ') AS protocolos
           FROM ({MOV}) m JOIN produtos p ON p.id = m.produto_id
           WHERE m.tipo = 'devolucao_estoque' AND m.ocorrido_em >= %(inicio)s AND m.ocorrido_em < %(fim_exclusivo)s
             AND {filial(m.filial_id)} AND {fornecedor(p.fornecedor_id)}
           GROUP BY p.sku, p.descricao, p.unidade ORDER BY p.sku"""),
    "requisicoes": Relatorio(
        "Requisições no período",
        [("protocolo", "Protocolo", "texto"), ("enviado_em", "Enviada em", "data"), ("filial", "Filial", "texto"),
         ("requisitante", "Requisitante", "texto"), ("status", "Status", "texto"),
         ("valor_solicitado", "Valor solicitado", "moeda"), ("valor_expedido", "Valor expedido", "moeda"),
         ("valor_em_loja", "Valor em loja", "moeda"), ("dias_ate_expedicao", "Dias até expedir", "qtd")],
        """SELECT s.protocolo, s.enviado_em, f.nome AS filial, u.nome AS requisitante,
                  CASE s.status WHEN 'aguardando_aprovacao_comercial' THEN 'Aguardando comercial'
                       WHEN 'aguardando_aprovacao_logistica' THEN 'Aguardando logística'
                       WHEN 'aguardando_expedicao' THEN 'Aguardando expedição' WHEN 'em_exposicao' THEN 'Em exposição'
                       WHEN 'devolucao_parcial' THEN 'Devolução parcial' WHEN 'encerrada' THEN 'Encerrada'
                       WHEN 'reprovada' THEN 'Reprovada' WHEN 'cancelada' THEN 'Cancelada' ELSE s.status::text END AS status,
                  (SELECT sum(i.quantidade_solicitada * i.custo_referencia) FROM itens_solicitacao i
                    WHERE i.solicitacao_id = s.id) AS valor_solicitado,
                  (SELECT sum(m.quantidade * m.custo_unitario) FROM ({MOV}) m
                    WHERE m.solicitacao_id = s.id AND m.tipo = 'expedicao') AS valor_expedido,
                  (SELECT sum(v.valor_em_loja) FROM vw_posicao_item v WHERE v.solicitacao_id = s.id) AS valor_em_loja,
                  round((extract(epoch FROM (SELECT min(m.ocorrido_em) FROM ({MOV}) m
                      WHERE m.solicitacao_id = s.id AND m.tipo = 'expedicao') - s.enviado_em) / 86400)::numeric, 1)
                    AS dias_ate_expedicao
           FROM solicitacoes s JOIN filiais f ON f.id = s.filial_id JOIN usuarios u ON u.id = s.requisitante_id
           WHERE s.status <> 'rascunho' AND s.enviado_em >= %(inicio)s AND s.enviado_em < %(fim_exclusivo)s
             AND {filial(s.filial_id)}
           ORDER BY s.enviado_em DESC LIMIT 5000"""),
    "aprovacoes_pendentes": Relatorio(
        "Aprovações pendentes",
        [("protocolo", "Protocolo", "texto"), ("etapa", "Etapa", "texto"), ("filial", "Filial", "texto"),
         ("requisitante", "Requisitante", "texto"), ("aguardando_desde", "Aguardando desde", "datahora"),
         ("dias", "Dias de espera", "qtd"), ("valor", "Valor solicitado", "moeda")],
        """SELECT s.protocolo,
                  CASE s.status WHEN 'aguardando_aprovacao_comercial' THEN 'Comercial' ELSE 'Logística' END AS etapa,
                  f.nome AS filial, u.nome AS requisitante,
                  COALESCE((SELECT decidido_em FROM aprovacoes a WHERE a.solicitacao_id = s.id AND a.etapa = 'comercial'),
                           s.enviado_em) AS aguardando_desde,
                  round((extract(epoch FROM now() - COALESCE((SELECT decidido_em FROM aprovacoes a
                      WHERE a.solicitacao_id = s.id AND a.etapa = 'comercial'), s.enviado_em)) / 86400)::numeric, 1) AS dias,
                  (SELECT sum(i.quantidade_solicitada * i.custo_referencia) FROM itens_solicitacao i
                    WHERE i.solicitacao_id = s.id) AS valor
           FROM solicitacoes s JOIN filiais f ON f.id = s.filial_id JOIN usuarios u ON u.id = s.requisitante_id
           WHERE s.status IN ('aguardando_aprovacao_comercial', 'aguardando_aprovacao_logistica')
             AND {filial(s.filial_id)}
           ORDER BY aguardando_desde""",
        usa_periodo=False),
}


def _montar_sql(rel: Relatorio, f: Filtro) -> str:
    return (rel.sql.replace("{MOV}", MOV_VALIDAS)
            .replace("{filial(v.filial_id)}", f.filial("v.filial_id"))
            .replace("{filial(m.filial_id)}", f.filial("m.filial_id"))
            .replace("{filial(s.filial_id)}", f.filial("s.filial_id"))
            .replace("{fornecedor(v.fornecedor_id)}", f.fornecedor("v.fornecedor_id"))
            .replace("{fornecedor(p.fornecedor_id)}", f.fornecedor("p.fornecedor_id")))


def _formatar_xlsx(rel: Relatorio, linhas: list[dict], f: Filtro) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = rel.titulo[:31]
    ws.append([rel.titulo])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Período: {f.inicio:%d/%m/%Y} a {f.fim:%d/%m/%Y}" if rel.usa_periodo else "Posição atual"])
    ws.append([])
    ws.append([c[1] for c in rel.colunas])
    for celula in ws[4]:
        celula.font = Font(bold=True)
    formatos = {"moeda": 'R$ #,##0.00', "qtd": '#,##0.###', "pct": '0.0%', "int": '0',
                "data": 'DD/MM/YYYY', "datahora": 'DD/MM/YYYY HH:MM'}
    for linha in linhas:
        valores = []
        for campo, _, tipo in rel.colunas:
            v = linha.get(campo)
            if tipo in ("moeda", "qtd", "pct") and v is not None:
                v = float(v)
            elif tipo in ("data", "datahora") and v is not None and getattr(v, "tzinfo", None):
                v = v.replace(tzinfo=None)
            valores.append(v)
        ws.append(valores)
        for celula, (_, _, tipo) in zip(ws[ws.max_row], rel.colunas):
            if tipo in formatos:
                celula.number_format = formatos[tipo]
    for n, (_, rotulo, tipo) in enumerate(rel.colunas, start=1):
        ws.column_dimensions[ws.cell(4, n).column_letter].width = max(12, min(45, len(rotulo) + 6 if tipo != "texto" else 28))
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def _formatar_csv(rel: Relatorio, linhas: list[dict]) -> bytes:
    saida = io.StringIO()
    w = csv.writer(saida, delimiter=";")
    w.writerow([c[1] for c in rel.colunas])
    for linha in linhas:
        valores = []
        for campo, _, tipo in rel.colunas:
            v = linha.get(campo)
            if v is None:
                v = ""
            elif tipo in ("moeda", "qtd", "pct"):
                v = str(v).replace(".", ",")
            elif tipo == "data":
                v = f"{v:%d/%m/%Y}"
            elif tipo == "datahora":
                v = f"{v:%d/%m/%Y %H:%M}"
            valores.append(v)
        w.writerow(valores)
    return ("﻿" + saida.getvalue()).encode("utf-8")  # BOM: o Excel abre com acentos


@router.get("/{nome}")
def relatorio(nome: str, formato: str = "json", f: Filtro = Depends(filtro), conn: Connection = Depends(get_conn)):
    rel = RELATORIOS.get(nome)
    if rel is None:
        raise HTTPException(404, "Relatório inexistente")
    linhas = conn.execute(_montar_sql(rel, f), f.params()).fetchall()
    if formato == "json":
        return {"titulo": rel.titulo, "usa_periodo": rel.usa_periodo,
                "colunas": [{"campo": c, "rotulo": r, "tipo": t} for c, r, t in rel.colunas], "linhas": linhas}
    if formato == "xlsx":
        conteudo, tipo, ext = _formatar_xlsx(rel, linhas, f), \
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    elif formato == "csv":
        conteudo, tipo, ext = _formatar_csv(rel, linhas), "text/csv; charset=utf-8", "csv"
    else:
        raise HTTPException(400, "Formato inválido")
    return StreamingResponse(io.BytesIO(conteudo), media_type=tipo,
                             headers={"Content-Disposition": f'attachment; filename="{nome}_{f.fim:%Y%m%d}.{ext}"'})
