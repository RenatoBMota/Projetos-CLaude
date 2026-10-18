"""Cadastros mestres: filiais, fornecedores, produtos e motivos.

Leitura liberada a qualquer usuário logado (os formulários precisam das
listas); gravação somente para administradores. Registros não são apagados,
apenas inativados, para preservar o histórico.
"""
import csv
import io
import unicodedata
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from psycopg import Connection, sql
from pydantic import BaseModel, Field

from ..db import auditar, get_conn
from ..seguranca import exigir_perfil, usuario_atual

router = APIRouter(prefix="/api/cadastros", tags=["cadastros"])
somente_admin = exigir_perfil("admin")


class FilialIn(BaseModel):
    codigo: str = Field(min_length=1)
    nome: str = Field(min_length=1)
    cnpj: str | None = None
    ativo: bool = True


class FornecedorIn(BaseModel):
    codigo: str = Field(min_length=1)
    razao_social: str = Field(min_length=1)
    cnpj: str | None = None
    contato: str | None = None
    ativo: bool = True


class ProdutoIn(BaseModel):
    sku: str = Field(min_length=1)
    descricao: str = Field(min_length=1)
    marca: str | None = None
    linha: str | None = None
    formato: str | None = None
    unidade: str = "PC"
    fornecedor_id: int | None = None
    custo_referencia: Decimal = Field(ge=0, default=Decimal(0))
    ativo: bool = True


class MotivoIn(BaseModel):
    tipo: Literal["descarte", "perda", "reprovacao", "cancelamento", "estorno"]
    descricao: str = Field(min_length=1)
    exige_observacao: bool = False
    ativo: bool = True


@dataclass(frozen=True)
class Entidade:
    tabela: str
    modelo: type[BaseModel]
    busca: tuple[str, ...]
    ordem: tuple[str, ...]


ENTIDADES = {
    "filiais": Entidade("filiais", FilialIn, ("codigo", "nome"), ("nome",)),
    "fornecedores": Entidade("fornecedores", FornecedorIn, ("codigo", "razao_social"), ("razao_social",)),
    "produtos": Entidade("produtos", ProdutoIn, ("sku", "descricao", "marca"), ("descricao",)),
    "motivos": Entidade("motivos", MotivoIn, ("descricao",), ("tipo", "descricao")),
}


def _entidade(nome: str) -> Entidade:
    if nome not in ENTIDADES:
        raise HTTPException(404, "Cadastro inexistente")
    return ENTIDADES[nome]


def _limpar(dados: BaseModel) -> dict:
    return {k: (v.strip() or None) if isinstance(v, str) else v for k, v in dados.model_dump().items()}


def _buscar(conn: Connection, ent: Entidade, registro_id: int) -> dict | None:
    return conn.execute(
        sql.SQL("SELECT * FROM {} WHERE id = %s").format(sql.Identifier(ent.tabela)), (registro_id,)
    ).fetchone()


# ---------------------------------------------------------------------------
# Importação de produtos (planilha exportada do ERP)
# Rotas fixas declaradas antes das genéricas /{entidade}.
# ---------------------------------------------------------------------------

COLUNAS_MODELO = ["sku", "descricao", "marca", "linha", "formato", "unidade",
                  "fornecedor_codigo", "fornecedor_nome", "custo"]

APELIDOS = {
    "codigo": "sku", "cod_produto": "sku", "codigo_produto": "sku",
    "produto": "descricao", "descricao_produto": "descricao",
    "un": "unidade", "und": "unidade",
    "cod_fornecedor": "fornecedor_codigo", "codigo_fornecedor": "fornecedor_codigo",
    "fornecedor": "fornecedor_nome", "razao_social": "fornecedor_nome",
    "custo_referencia": "custo", "custo_medio": "custo", "custo_unitario": "custo",
}


def _normalizar_coluna(nome) -> str:
    texto = unicodedata.normalize("NFKD", str(nome or "")).encode("ascii", "ignore").decode()
    chave = "_".join(texto.strip().lower().replace("-", " ").split())
    return APELIDOS.get(chave, chave)


def _texto(valor) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip() or None


def _custo(valor) -> Decimal:
    if valor is None or valor == "":
        return Decimal(0)
    if isinstance(valor, (int, float)):
        return Decimal(str(round(valor, 4)))
    texto = str(valor).replace("R$", "").replace(" ", "")
    if "," in texto:  # formato brasileiro: 1.234,56
        texto = texto.replace(".", "").replace(",", ".")
    custo = Decimal(texto)
    if custo < 0:
        raise InvalidOperation
    return custo


def _ler_planilha(arquivo: UploadFile) -> list[dict]:
    conteudo = arquivo.file.read()
    nome = (arquivo.filename or "").lower()
    if nome.endswith(".xlsx"):
        planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).active
        linhas = list(planilha.iter_rows(values_only=True))
    elif nome.endswith(".csv"):
        texto = conteudo.decode("utf-8-sig", errors="replace")
        separador = ";" if texto.split("\n", 1)[0].count(";") >= texto.split("\n", 1)[0].count(",") else ","
        linhas = list(csv.reader(io.StringIO(texto), delimiter=separador))
    else:
        raise HTTPException(400, "Envie um arquivo .xlsx ou .csv")
    if not linhas:
        raise HTTPException(400, "Planilha vazia")
    cabecalho = [_normalizar_coluna(c) for c in linhas[0]]
    faltando = {"sku", "descricao"} - set(cabecalho)
    if faltando:
        raise HTTPException(400, f"Colunas obrigatórias ausentes: {', '.join(sorted(faltando))}")
    return [dict(zip(cabecalho, linha)) for linha in linhas[1:] if any(c not in (None, "") for c in linha)]


@router.get("/produtos/modelo")
def modelo_importacao(_: dict = Depends(usuario_atual)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Produtos"
    ws.append(COLUNAS_MODELO)
    ws.append(["PORC-6060-01", "Porcelanato Acetinado 60x60 Cinza", "Marca X", "Linha Y",
               "60x60", "M2", "F001", "Cerâmica Exemplo Ltda", "45,90"])
    saida = io.BytesIO()
    wb.save(saida)
    saida.seek(0)
    return StreamingResponse(
        saida, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="modelo_produtos.xlsx"'},
    )


@router.post("/produtos/importar")
def importar_produtos(arquivo: UploadFile, admin: dict = Depends(somente_admin),
                      conn: Connection = Depends(get_conn)):
    """Insere ou atualiza produtos pelo SKU. Tudo ou nada: com qualquer erro, nada é gravado."""
    linhas = _ler_planilha(arquivo)
    erros, validas, skus = [], [], set()
    for n, linha in enumerate(linhas, start=2):
        sku, descricao = _texto(linha.get("sku")), _texto(linha.get("descricao"))
        if not sku or not descricao:
            erros.append({"linha": n, "mensagem": "SKU e descrição são obrigatórios"})
            continue
        if sku in skus:
            erros.append({"linha": n, "mensagem": f"SKU {sku} repetido na planilha"})
            continue
        skus.add(sku)
        try:
            custo = _custo(linha.get("custo"))
        except (InvalidOperation, ValueError):
            erros.append({"linha": n, "mensagem": f"Custo inválido: {linha.get('custo')}"})
            continue
        validas.append((n, {
            "sku": sku, "descricao": descricao,
            "marca": _texto(linha.get("marca")), "linha": _texto(linha.get("linha")),
            "formato": _texto(linha.get("formato")), "unidade": _texto(linha.get("unidade")) or "PC",
            "custo_referencia": custo,
            "fornecedor_codigo": _texto(linha.get("fornecedor_codigo")),
            "fornecedor_nome": _texto(linha.get("fornecedor_nome")),
        }))

    fornecedores = {r["codigo"]: r["id"] for r in conn.execute("SELECT id, codigo FROM fornecedores")}
    novos_fornecedores = {}
    for n, p in validas:
        cod = p["fornecedor_codigo"]
        if cod and cod not in fornecedores and cod not in novos_fornecedores:
            if not p["fornecedor_nome"]:
                erros.append({"linha": n, "mensagem": f"Fornecedor {cod} não cadastrado; informe fornecedor_nome"})
            else:
                novos_fornecedores[cod] = p["fornecedor_nome"]

    if erros:
        return {"importado": False, "erros": sorted(erros, key=lambda e: e["linha"])}

    for cod, nome in novos_fornecedores.items():
        fornecedores[cod] = conn.execute(
            "INSERT INTO fornecedores (codigo, razao_social) VALUES (%s, %s) RETURNING id", (cod, nome)
        ).fetchone()["id"]

    inseridos = atualizados = 0
    for _, p in validas:
        r = conn.execute(
            """INSERT INTO produtos (sku, descricao, marca, linha, formato, unidade, fornecedor_id, custo_referencia)
               VALUES (%(sku)s, %(descricao)s, %(marca)s, %(linha)s, %(formato)s, %(unidade)s, %(fornecedor_id)s, %(custo_referencia)s)
               ON CONFLICT (sku) DO UPDATE SET
                 descricao = EXCLUDED.descricao, marca = EXCLUDED.marca, linha = EXCLUDED.linha,
                 formato = EXCLUDED.formato, unidade = EXCLUDED.unidade,
                 fornecedor_id = COALESCE(EXCLUDED.fornecedor_id, produtos.fornecedor_id),
                 custo_referencia = EXCLUDED.custo_referencia
               RETURNING (xmax = 0) AS inserido""",
            {**p, "fornecedor_id": fornecedores.get(p["fornecedor_codigo"])},
        ).fetchone()
        if r["inserido"]:
            inseridos += 1
        else:
            atualizados += 1

    auditar(conn, admin["id"], "produtos", 0, "importar",
            depois={"arquivo": arquivo.filename, "inseridos": inseridos, "atualizados": atualizados,
                    "fornecedores_criados": len(novos_fornecedores)})
    return {"importado": True, "inseridos": inseridos, "atualizados": atualizados,
            "fornecedores_criados": len(novos_fornecedores), "erros": []}


# ---------------------------------------------------------------------------
# CRUD genérico
# ---------------------------------------------------------------------------

@router.get("/{entidade}")
def listar(entidade: str, busca: str = "", ativos: bool | None = None, pagina: int = 1,
           limite: int = 50, _: dict = Depends(usuario_atual), conn: Connection = Depends(get_conn)):
    ent = _entidade(entidade)
    limite = max(1, min(limite, 500))
    filtros, params = [], []
    if busca.strip():
        filtros.append(sql.SQL("({})").format(sql.SQL(" OR ").join(
            sql.SQL("unaccent(t.{}) ILIKE unaccent(%s)").format(sql.Identifier(c)) for c in ent.busca)))
        params += [f"%{busca.strip()}%"] * len(ent.busca)
    if ativos is not None:
        filtros.append(sql.SQL("t.ativo = %s"))
        params.append(ativos)
    onde = sql.SQL("WHERE ") + sql.SQL(" AND ").join(filtros) if filtros else sql.SQL("")
    extra, juncao = sql.SQL(""), sql.SQL("")
    if entidade == "produtos":
        extra = sql.SQL(", f.razao_social AS fornecedor_nome, f.codigo AS fornecedor_codigo")
        juncao = sql.SQL("LEFT JOIN fornecedores f ON f.id = t.fornecedor_id")
    consulta = sql.SQL(
        "SELECT t.*{extra}, count(*) OVER () AS _total FROM {tabela} t {juncao} {onde} "
        "ORDER BY {ordem} LIMIT %s OFFSET %s"
    ).format(extra=extra, tabela=sql.Identifier(ent.tabela), juncao=juncao, onde=onde,
             ordem=sql.SQL(", ").join(sql.SQL("t.{}").format(sql.Identifier(c)) for c in ent.ordem))
    linhas = conn.execute(consulta, params + [limite, (max(pagina, 1) - 1) * limite]).fetchall()
    total = linhas[0]["_total"] if linhas else 0
    for linha in linhas:
        del linha["_total"]
    return {"itens": linhas, "total": total, "pagina": pagina, "limite": limite}


@router.post("/{entidade}", status_code=201)
def criar(entidade: str, corpo: dict, admin: dict = Depends(somente_admin),
          conn: Connection = Depends(get_conn)):
    ent = _entidade(entidade)
    dados = _limpar(ent.modelo.model_validate(corpo))
    colunas = list(dados)
    novo = conn.execute(
        sql.SQL("INSERT INTO {} ({}) VALUES ({}) RETURNING *").format(
            sql.Identifier(ent.tabela), sql.SQL(", ").join(map(sql.Identifier, colunas)),
            sql.SQL(", ").join(sql.Placeholder(c) for c in colunas)),
        dados,
    ).fetchone()
    auditar(conn, admin["id"], ent.tabela, novo["id"], "criar", depois=novo)
    return novo


@router.put("/{entidade}/{registro_id}")
def alterar(entidade: str, registro_id: int, corpo: dict, admin: dict = Depends(somente_admin),
            conn: Connection = Depends(get_conn)):
    ent = _entidade(entidade)
    dados = _limpar(ent.modelo.model_validate(corpo))
    antes = _buscar(conn, ent, registro_id)
    if antes is None:
        raise HTTPException(404, "Registro não encontrado")
    depois = conn.execute(
        sql.SQL("UPDATE {} SET {} WHERE id = %(id)s RETURNING *").format(
            sql.Identifier(ent.tabela),
            sql.SQL(", ").join(sql.SQL("{} = {}").format(sql.Identifier(c), sql.Placeholder(c)) for c in dados)),
        {**dados, "id": registro_id},
    ).fetchone()
    auditar(conn, admin["id"], ent.tabela, registro_id, "alterar", antes, depois)
    return depois
