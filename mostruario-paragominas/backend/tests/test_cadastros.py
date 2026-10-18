import io

from openpyxl import Workbook, load_workbook


def test_crud_fornecedor_e_produto(cliente, admin):
    f = cliente.post("/api/cadastros/fornecedores", headers=admin,
                     json={"codigo": "F01", "razao_social": "Cerâmica A"})
    assert f.status_code == 201
    p = cliente.post("/api/cadastros/produtos", headers=admin, json={
        "sku": "P1", "descricao": "Porcelanato 60x60", "fornecedor_id": f.json()["id"],
        "custo_referencia": "45.9"})
    assert p.status_code == 201, p.text

    lista = cliente.get("/api/cadastros/produtos?busca=porcel", headers=admin).json()
    assert lista["total"] == 1
    assert lista["itens"][0]["fornecedor_nome"] == "Cerâmica A"
    assert float(lista["itens"][0]["custo_referencia"]) == 45.9

    alt = cliente.put(f"/api/cadastros/produtos/{p.json()['id']}", headers=admin, json={
        "sku": "P1", "descricao": "Porcelanato 60x60", "custo_referencia": 50, "ativo": False})
    assert alt.status_code == 200 and alt.json()["ativo"] is False
    assert cliente.get("/api/cadastros/produtos?ativos=true", headers=admin).json()["total"] == 0


def test_validacoes(cliente, admin):
    assert cliente.post("/api/cadastros/produtos", headers=admin,
                        json={"sku": "", "descricao": "x"}).status_code == 422
    assert cliente.post("/api/cadastros/produtos", headers=admin,
                        json={"sku": "A", "descricao": "x", "custo_referencia": -1}).status_code == 422
    assert cliente.post("/api/cadastros/produtos", headers=admin,
                        json={"sku": "A", "descricao": "x", "fornecedor_id": 999}).status_code == 400
    cliente.post("/api/cadastros/filiais", headers=admin, json={"codigo": "PGM", "nome": "Paragominas"})
    assert cliente.post("/api/cadastros/filiais", headers=admin,
                        json={"codigo": "PGM", "nome": "Outra"}).status_code == 409
    assert cliente.get("/api/cadastros/inexistente", headers=admin).status_code == 404


def test_motivos_iniciais(cliente, admin):
    itens = cliente.get("/api/cadastros/motivos?limite=500", headers=admin).json()["itens"]
    descarte = [m["descricao"] for m in itens if m["tipo"] == "descarte"]
    assert "Quebrada / trincada" in descarte


def _xlsx(linhas):
    wb = Workbook()
    for linha in linhas:
        wb.active.append(linha)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_importar_planilha(cliente, admin):
    modelo = cliente.get("/api/cadastros/produtos/modelo", headers=admin)
    cabecalho = [c.value for c in next(load_workbook(io.BytesIO(modelo.content)).active.iter_rows())]

    arquivo = _xlsx([cabecalho,
                     ["P1", "Porcelanato A", "M", "L", "60x60", "M2", "F01", "Cerâmica A", "1.234,56"],
                     ["P2", "Porcelanato B", None, None, None, None, "F01", None, 30.5]])
    r = cliente.post("/api/cadastros/produtos/importar", headers=admin,
                     files={"arquivo": ("produtos.xlsx", arquivo)})
    assert r.json() == {"importado": True, "inseridos": 2, "atualizados": 0,
                        "fornecedores_criados": 1, "erros": []}

    # reimportação atualiza pelo SKU; CSV com ; e cabeçalhos com acento
    csv = "Código;Descrição;Custo\nP1;Porcelanato A novo;10,00\n".encode()
    r = cliente.post("/api/cadastros/produtos/importar", headers=admin,
                     files={"arquivo": ("produtos.csv", csv)})
    assert r.json()["atualizados"] == 1
    p1 = cliente.get("/api/cadastros/produtos?busca=P1", headers=admin).json()["itens"][0]
    assert p1["descricao"] == "Porcelanato A novo" and float(p1["custo_referencia"]) == 10
    assert p1["fornecedor_nome"] == "Cerâmica A"  # mantém fornecedor quando a coluna vem vazia


def test_importacao_com_erro_nao_grava_nada(cliente, admin):
    arquivo = _xlsx([["sku", "descricao", "custo", "fornecedor_codigo"],
                     ["P1", "Ok", "10", None],
                     ["P2", None, "10", None],
                     ["P3", "Custo ruim", "abc", None],
                     ["P4", "Fornecedor novo sem nome", "1", "F99"]])
    r = cliente.post("/api/cadastros/produtos/importar", headers=admin,
                     files={"arquivo": ("p.xlsx", arquivo)}).json()
    assert r["importado"] is False
    assert [e["linha"] for e in r["erros"]] == [3, 4, 5]
    assert cliente.get("/api/cadastros/produtos", headers=admin).json()["total"] == 0


def test_importacao_formato_invalido(cliente, admin):
    r = cliente.post("/api/cadastros/produtos/importar", headers=admin,
                     files={"arquivo": ("p.txt", b"x")})
    assert r.status_code == 400
    r = cliente.post("/api/cadastros/produtos/importar", headers=admin,
                     files={"arquivo": ("p.csv", b"nome;valor\na;1\n")})
    assert r.status_code == 400 and "sku" in r.json()["detail"]
