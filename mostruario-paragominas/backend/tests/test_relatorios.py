import io
import uuid

import pytest
from openpyxl import load_workbook

from test_fluxo import _acao, _aprovar_tudo, _nova, cenario  # noqa: F401
from app.rotas.relatorios import RELATORIOS


@pytest.fixture
def movimentado(cenario):  # noqa: F811
    """Requisição expedida e parcialmente devolvida (estoque, descarte e perda)."""
    k = cenario
    s = _aprovar_tudo(k, _nova(k))
    i1, i2 = s["itens"][0]["item_id"], s["itens"][1]["item_id"]
    _acao(k, "dep", f"/api/solicitacoes/{s['id']}/expedir",
          {"chave": str(uuid.uuid4()), "itens": [{"item_id": i1, "quantidade": 10}, {"item_id": i2, "quantidade": 4}]})
    m = k["motivos"]
    r = _acao(k, "dep", f"/api/solicitacoes/{s['id']}/devolver", {"chave": str(uuid.uuid4()), "lancamentos": [
        {"item_id": i1, "destino": "estoque", "quantidade": 6},
        {"item_id": i1, "destino": "descarte", "quantidade": 2, "motivo_id": m[("descarte", "Riscada / manchada")]},
        {"item_id": i2, "destino": "perda", "quantidade": 2, "motivo_id": m[("perda", "Extraviada na loja")]}]})
    assert r.status_code == 200, r.text
    return k


def test_indicadores(movimentado):
    k = movimentado
    d = k["c"].get("/api/relatorios/indicadores", headers=k["ges"]).json()
    assert d["solicitacoes"] == 1
    assert float(d["valor_solicitado"]) == 10 * 50 + 4 * 20
    assert float(d["valor_expedido"]) == 580
    assert [float(d[c]) for c in ("qtd_estoque", "qtd_descarte", "qtd_perda")] == [6, 2, 2]
    assert d["taxa_reaproveitamento"] == 0.6 and d["taxa_descarte"] == 0.2 and d["taxa_perda"] == 0.2
    assert float(d["em_exposicao"]["valor"]) == 2 * 50 + 2 * 20
    assert d["taxa_aprovacao_comercial"] == 1.0 and d["taxa_aprovacao_logistica"] == 1.0
    assert d["horas_aprovacao_comercial"] is not None


def test_todos_os_relatorios_e_exportacoes(movimentado):
    k = movimentado
    for nome in RELATORIOS:
        r = k["c"].get(f"/api/relatorios/{nome}", headers=k["ges"])
        assert r.status_code == 200, (nome, r.text)
        dados = r.json()
        assert dados["colunas"] and isinstance(dados["linhas"], list)
        x = k["c"].get(f"/api/relatorios/{nome}?formato=xlsx", headers=k["ges"])
        assert x.status_code == 200, (nome, x.text)
        ws = load_workbook(io.BytesIO(x.content)).active
        assert ws.max_row == 4 + len(dados["linhas"])
        c = k["c"].get(f"/api/relatorios/{nome}?formato=csv", headers=k["ges"])
        assert c.content.startswith("﻿".encode())

    erp = k["c"].get("/api/relatorios/lancar_erp", headers=k["dep"]).json()["linhas"]
    assert [(l["sku"], float(l["quantidade"])) for l in erp] == [("P1", 6)]
    motivos = k["c"].get("/api/relatorios/motivos", headers=k["ges"]).json()["linhas"]
    assert {(l["destino"], l["motivo"]) for l in motivos} == {("Descarte", "Riscada / manchada"),
                                                             ("Perda", "Extraviada na loja")}
    por_filial = k["c"].get("/api/relatorios/exposicao_filial", headers=k["ges"]).json()["linhas"]
    assert [(l["filial"], float(l["valor"])) for l in por_filial] == [("Paragominas", 140)]


def test_permissoes_e_periodo(movimentado):
    k = movimentado
    assert k["c"].get("/api/relatorios/indicadores", headers=k["req"]).status_code == 403
    assert k["c"].get("/api/relatorios/inexistente", headers=k["ges"]).status_code == 404
    assert k["c"].get("/api/relatorios/movimentacoes?formato=pdf", headers=k["ges"]).status_code == 400
    r = k["c"].get("/api/relatorios/indicadores?inicio=2020-01-01&fim=2020-01-31", headers=k["ges"]).json()
    assert r["solicitacoes"] == 0 and float(r["qtd_expedida"]) == 0
    assert k["c"].get("/api/relatorios/indicadores?inicio=2026-02-01&fim=2026-01-01",
                      headers=k["ges"]).status_code == 400
    # aprovador de Belém não enxerga Paragominas
    h = k["bel_req"]
    assert k["c"].get("/api/relatorios/exposicao_filial", headers=h).json()["linhas"] == []


PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
       b"\x00\x00\x00\rIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82")


def test_anexos(cenario, tmp_path, monkeypatch):  # noqa: F811
    from app import config
    monkeypatch.setattr(config, "ANEXOS_DIR", str(tmp_path))
    k = cenario
    s = _nova(k)
    url = f"/api/solicitacoes/{s['id']}/anexos"
    r = k["c"].post(url, headers=k["dep"], files={"arquivo": ("foto.png", PNG, "image/png")},
                    data={"descricao": "Comprovante de coleta"})
    assert r.status_code == 201, r.text
    lista = k["c"].get(url, headers=k["req"]).json()
    assert lista[0]["descricao"] == "Comprovante de coleta" and lista[0]["usuario_nome"] == "Dep"
    baixado = k["c"].get(f"/api/anexos/{lista[0]['id']}", headers=k["req"])
    assert baixado.content == PNG and baixado.headers["content-type"] == "image/png"

    assert k["c"].post(url, headers=k["dep"], files={"arquivo": ("x.exe", b"MZ", "application/x-msdownload")}).status_code == 400
    # fora da abrangência
    assert k["c"].get(url, headers=k["bel_req"]).status_code == 404
    assert k["c"].get(f"/api/anexos/{lista[0]['id']}", headers=k["bel_req"]).status_code == 404
    historico = k["c"].get(f"/api/solicitacoes/{s['id']}", headers=k["req"]).json()["historico"]
    assert any(h["acao"] == "anexar" for h in historico)


def test_painel_gerencial(movimentado):
    k = movimentado
    d = k["c"].get("/api/relatorios/gerencial", headers=k["ges"]).json()
    assert len(d["mensal"]) >= 3  # 90 dias cobrem ao menos 3 meses
    atual = d["mensal"][-1]
    assert atual["requisicoes"] == 1 and float(atual["expedido"]) == 580
    assert [float(atual[c]) for c in ("estoque", "descarte", "perda")] == [300, 100, 40]
    assert [(x["filial"], float(x["estoque"]), float(x["descarte"]), float(x["perda"]))
            for x in d["destino_filial"]] == [("Paragominas", 6, 2, 2)]
    assert [a["faixa"] for a in d["aging"]][0] == "0–30 dias"
    assert d["aging"][0]["itens"] == 2 and float(d["aging"][0]["valor"]) == 140
    assert [p["sku"] for p in d["top_produtos"]] == ["P1", "P2"]
    assert d["top_requisitantes"][0]["requisitante"] == "Req"
    assert d["status"] == [{"status": "devolucao_parcial", "n": 1}]
    assert k["c"].get("/api/relatorios/gerencial", headers=k["req"]).status_code == 403
